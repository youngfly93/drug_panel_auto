# ruff: noqa: E402, I001
"""肺癌329+PD-L1 受控试运行冒烟与安全边界。

默认模板必须无病例硬编码、可渲染；PD-L1 必须逐病例提供来源，
批量生成和未经人工复核的下载必须保持 fail-closed。
"""

import io
import hashlib
import re
import sys
import zipfile
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

_PANEL_DIR = ROOT / "panels" / "lung_329_pdl1"


def _load_panel_spec() -> dict:
    return yaml.safe_load((_PANEL_DIR / "panel.yaml").read_text(encoding="utf-8"))


def _resolve_default_template(spec: dict) -> Path:
    """跟随 panel.yaml 的 default_template 解析模板文件，避免硬编码漂移。

    Codex P2：default 切到 v1 后，golden case 不能还只渲染/扫 v0——否则
    生产实际用的模板从不被本包自己的 golden case 渲染或 PII 扫描。
    """
    default_id = spec["default_template"]
    file_rel = next(t["file"] for t in spec["templates"] if t.get("id") == default_id)
    return _PANEL_DIR / file_rel


_PANEL_SPEC = _load_panel_spec()
TEMPLATE = _resolve_default_template(_PANEL_SPEC)
# 模板的循环集合（{%tr for ... %}）。冒烟渲染给空列表，断言零行不崩。
_REQUIRED_LISTS = list(
    (_PANEL_SPEC.get("template_contract") or {}).get("required_lists") or []
)

_MVP_SCALARS = [
    "patient_name",
    "sample_id",
    "gender",
    "age",
    "clinical_diagnosis",
    "sample_type",
    "sampling_method",
    "sample_site",
    "tmb_summary",
    "msi_summary",
    "pdl1_tps",
    "pdl1_cps",
    "pdl1_result",
    "immune_positive_result",
    "immune_negative_result",
    "immune_hyperprogression_result",
]


def test_lung329_template_exists():
    assert TEMPLATE.exists(), f"肺329模板缺失: {TEMPLATE}"
    assert TEMPLATE.stat().st_size > 10_000, "肺329模板文件异常小"
    registered = {
        entry["file"]
        for entry in (_PANEL_SPEC.get("templates") or [])
    }
    assert registered == {
        "templates/lung_329_pdl1_golden_template_v2.docx"
    }
    assert not (
        _PANEL_DIR / "templates" / "lung_329_pdl1_golden_template_v0.docx"
    ).exists()
    assert not (
        _PANEL_DIR / "templates" / "lung_329_pdl1_golden_template_v1.docx"
    ).exists()


def test_lung329_template_is_pii_clean():
    """通用扫描器必须找不到病例变异、日期、丰度或调试残留。"""
    from scripts.scan_hardcoded_literals import scan_docx

    result = scan_docx(TEMPLATE, tokens=[])
    assert result.matches == []


def test_lung329_template_has_no_orphan_or_source_case_media():
    """历史病例图片及已删除章节的孤儿媒体不得藏在DOCX压缩包内。"""
    with zipfile.ZipFile(TEMPLATE) as archive:
        names = set(archive.namelist())
        media_parts = {name for name in names if name.startswith("word/media/")}
        relationship_targets = {
            "word/media/" + target.decode("utf-8")
            for name in names
            if name.endswith(".rels")
            for target in re.findall(
                rb'Target="(?:\.\./)?media/([^"]+)"',
                archive.read(name),
            )
        }

    assert media_parts == relationship_targets
    assert len(media_parts) == 9
    assert "word/media/image6.jpeg" not in media_parts


def test_lung329_template_gene_universe_matches_governed_coverage():
    from docx import Document

    coverage = yaml.safe_load(
        (_PANEL_DIR / "rules" / "knowledge_coverage.yaml").read_text(
            encoding="utf-8"
        )
    )
    expected = coverage["reportable_genes"]
    actual_hash = hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()
    document = Document(TEMPLATE)
    gene_table = next(
        table
        for table in document.tables
        if "Gene List for MLseq" in table.cell(0, 0).text
    )
    actual = [
        cell.text.strip().upper()
        for row in gene_table.rows[1:]
        for cell in row.cells
        if cell.text.strip()
    ]

    assert actual_hash == coverage["source"]["template_sha256"]
    assert len(actual) == 329
    assert len(set(actual)) == 329
    assert actual == expected


def test_lung329_panel_registers():
    from reportgen.core.enhancer_registry import get_panel_registry

    reg = get_panel_registry()
    assert reg.get("lung_329_pdl1") is not None, "lung_329_pdl1 未注册"


def test_lung329_template_renders_with_scalars():
    """docxtpl 用哨兵标量 context 渲染，且绑定值确实进入文档。"""
    from docxtpl import DocxTemplate

    tpl = DocxTemplate(str(TEMPLATE))
    # 用不含尖括号的哨兵值（避免被 XML 去标签正则误删）
    ctx = {k: f"SENTINEL_{k}_VAL" for k in _MVP_SCALARS}
    ctx.update({k: [] for k in _REQUIRED_LISTS})  # 循环集合给空列表 → 渲染零行不报错
    tpl.render(ctx)

    buf = io.BytesIO()
    tpl.save(buf)
    buf.seek(0)
    with zipfile.ZipFile(buf) as z:
        doc = z.read("word/document.xml").decode("utf-8", "ignore")
    visible = re.sub(r"<[^>]+>", "", doc)
    assert "SENTINEL_patient_name_VAL" in visible
    assert "SENTINEL_pdl1_result_VAL" in visible


def test_lung329_pdl1_form_requires_case_specific_provenance():
    from app.services.clinical_info_service import get_clinical_form_schema

    schema = get_clinical_form_schema("lung_329_pdl1")
    fields = {field.key: field for group in schema.groups for field in group.fields}
    required = {
        "pdl1_tps",
        "pdl1_cps",
        "pdl1_result",
        "pdl1_assay_profile_id",
        "pdl1_source_record_id",
        "pdl1_source_record_date",
        "pdl1_specimen_id",
        "pdl1_image_disposition",
    }

    assert required <= set(fields)
    assert all(fields[key].required for key in required)
    assert fields["pdl1_assay_profile_id"].ui.options == [
        "legacy_unspecified_ihc_transcription_v1"
    ]
    assert fields["pdl1_image_disposition"].ui.options == [
        "无病例专属图像（报告不展示）"
    ]


def test_lung329_shared_batch_is_blocked():
    from app.api.batch import _batch_generation_policy_error

    error = _batch_generation_policy_error("lung_329_pdl1")
    assert error is not None
    assert "逐病例" in error
    assert "串用" in error


def test_lung329_explicit_empty_immune_categories_do_not_fall_back(tmp_path):
    from reportgen.core.template_bridge_358 import (
        _build_nccn_and_immune_fields,
        load_panel_config,
    )
    from reportgen.models.excel_data import ExcelDataSource
    from reportgen.models.report_data import ReportData
    from reportgen.panels.loader import load_panel_package

    package = load_panel_package("lung_329_pdl1", project_root=ROOT)
    panel_config = load_panel_config(
        panel_package=package,
    )
    report_data = ReportData()
    source = tmp_path / "synthetic-lung329.xlsx"
    source.touch()
    _build_nccn_and_immune_fields(
        report_data,
        [],
        ExcelDataSource(file_path=str(source)),
        panel_config=panel_config,
    )

    assert panel_config.declared_immune_categories == {
        "positive",
        "negative",
        "hyperprogression",
    }
    assert report_data.get_table("immune_positive_results") == []
    assert report_data.get_table("immune_negative_results") == []
    assert report_data.get_table("immune_hyperprogression_results") == []
